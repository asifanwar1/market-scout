import csv
import io
import zipfile
from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

SNIFF_CHARACTERS = 8192
SHEET_STATE_VISIBLE = "visible"


_MAX_UNPACKED_BYTES = 200 * 1024 * 1024
_CSV_ENCODINGS = ("utf-8-sig", "cp1252")
_DELIMITERS = ",;\t|"


class ImportFileError(Exception):
    pass


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]


def read_table(filename: str, content: bytes, max_rows: int) -> Table:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        raw = _read_csv(content)
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        raw = _read_xlsx(content, max_rows + 2)
    elif name.endswith(".xls"):
        raise ImportFileError("Old .xls files can't be read. Open it in Excel and save it as .xlsx or .csv.")
    else:
        raise ImportFileError("Upload a .csv or .xlsx file.")

    while raw and not any(cell.strip() for cell in raw[0]):
        raw.pop(0)
    if not raw:
        raise ImportFileError("The file is empty.")

    headers = _unique_headers(raw[0])
    width = len(headers)
    rows = [
        [(row[i] if i < len(row) else "").strip() for i in range(width)]
        for row in raw[1:]
        if any(cell.strip() for cell in row)
    ]
    if not rows:
        raise ImportFileError("The file has a header row but no data rows.")
    if len(rows) > max_rows:
        raise ImportFileError(f"The file has {len(rows):,} rows; the limit is {max_rows:,}. Split it into smaller files.")
    return Table(headers=headers, rows=rows)


def _unique_headers(first_row: list[str]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, cell in enumerate(first_row, start=1):
        title = " ".join(cell.split()) or f"Column {index}"
        seen[title.lower()] = seen.get(title.lower(), 0) + 1
        headers.append(title if seen[title.lower()] == 1 else f"{title} ({seen[title.lower()]})")
    return headers


def _read_csv(content: bytes) -> list[list[str]]:
    text = None
    for encoding in _CSV_ENCODINGS:
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ImportFileError("The file's text encoding isn't recognised. Save it as CSV (UTF-8) and try again.")

    sample = text[:SNIFF_CHARACTERS]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=_DELIMITERS)
    except csv.Error:
        dialect = csv.excel
    try:
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    except csv.Error as exc:
        raise ImportFileError(f"The CSV file is malformed: {exc}") from None


def _read_xlsx(content: bytes, limit: int) -> list[list[str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > _MAX_UNPACKED_BYTES:
                raise ImportFileError("This Excel file is too large once unpacked.")
    except zipfile.BadZipFile:
        raise ImportFileError("This isn't a valid .xlsx file.") from None

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise ImportFileError(f"This Excel file can't be read ({exc.__class__.__name__}).") from None

    try:
        rows: list[list[str]] = []
        sheet = next((s for s in workbook.worksheets if s.sheet_state == SHEET_STATE_VISIBLE), None)
        if sheet is None:
            raise ImportFileError("The workbook has no visible sheet.")
        for row in sheet.iter_rows(values_only=True):
            rows.append([_cell_text(value) for value in row])
            if len(rows) >= limit:
                break
        return rows
    finally:
        workbook.close()


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
